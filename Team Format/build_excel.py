"""Build teamFormatAnalysis.xlsx from team_format_model results.

Layout: one band per format (rows: Unchanged, Proposal 1, Proposal 1a); two blocks side by side per band
(left: Proposal 2 off, right: Proposal 2 on). Each block holds the three league tables, the participation tables
below them (3-colour scale on the participation counts) and one impact comment under each table.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as col

from team_format_model import run_all, LEGS

OUT = Path(__file__).resolve().parent / "teamFormatAnalysis.xlsx"
PROPOSALS = ["unchanged", "P1", "P1a"]
PNAME = {"unchanged": "Unchanged format", "P1": "Proposal 1", "P1a": "Proposal 1a"}
SUBTITLE = {  # minimal team-structure subtitle shown under each league table title
    "unchanged": {"ENCT": "Best 4 non-compound archers (novices included)", "NNCT": "Best 3 non-compound novices",
                  "CT": "Best 2 compound archers"},
    "P1": {"ENCT": "Best 3 non-compound archers (novices included, no compound)", "NNCT": "Best 3 non-compound novices",
           "UST": "Best 3 barebow/longbow archers (novices included)"},
    "P1a": {"ST": "Best 3 recurve + max 1 compound (L7-8 n/c)", "NNCT": "Best 3 non-compound novices (L7-8 n/c)",
            "UST": "Best 3 barebow/longbow archers (L7-8 n/c)"},
}
GENDER_SUB = "; needs >=1 Female and >=1 Open"
ROW_HEADINGS = {"unchanged": "UNCHANGED FORMAT", "P1": "PROPOSAL 1", "P1a": "PROPOSAL 1a"}
COL_HEADINGS = {False: "PROPOSAL 2 OFF: no gender requirement",
                True: "PROPOSAL 2 ON: each team needs at least 1 Female and 1 Open archer"}
HOW_TO_READ = (
    "How to read: League tables - Pos is the club's league position (ranked on league points, ties broken on aggregate). "
    "L1-L8 give the team's rank in each leg (1 = best); '-' means the club fielded no team that leg and 'n/c' means the leg "
    "does not count under that format. LP = league points (10 for 1st, 9 for 2nd, and so on down); Agg = aggregate, the sum "
    "of the team scores over the counted legs. Participation tables - number of archers whose scores counted towards the "
    "club's team in each leg (Total = sum over legs), shaded red (fewest) to green (most) within each table. "
    "Impact - one comment under each league table: leader, change in overall (league) rank versus a stated reference, and "
    "any gender shortfall. Shortfall entries read: club, leg (archers counted/team size, change in team score, change in leg "
    "rank; LOWER = ranked lower than without Proposal 2). See the 'Explanation' sheet for the formats and rules.")
REF_TEAM = {"ENCT": "ENCT", "ST": "ENCT", "NNCT": "NNCT", "CT": None, "UST": None}  # gender-off comparison vs unchanged

HEAD = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="999999")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BAND, BAND0, BLOCK_OFFSET, TABLE_STRIDE = 30, 6, 42, 13  # rows per band, first band row, cols between blocks/tables


def pos_changes(new: pd.Series, old: pd.Series) -> str:
    """Describe league position changes between two Pos series (index = club). Returns 'unchanged' or a list."""
    mv = [f"{c} {old[c]}->{new[c]}" for c in new.index if new[c] != old[c]]
    return ", ".join(mv) if mv else "unchanged"


def table_comment(p: str, g: bool, team: str, R: dict) -> str:
    """Impact comment for one league table, in a fixed three-line format.

    Line 1: leader and margin. Line 2: change in overall (league) rank versus a reference table.
    Line 3: gender shortfalls (single-gender pool => team below full size), with score and leg-rank effect.
    Args: p proposal key, g gender flag, team team name, R all results. Returns: str.
    """
    r = R[(p, g)][team]
    tb = r["table"]
    gap = tb.LP.iloc[0] - tb.LP.iloc[1]
    lead = (f"Leader: {tb.index[0]} ({tb.LP.iloc[0]:.0f} LP, +{gap:.0f} on {tb.index[1]})." if gap
            else f"Leader: {tb.index[0]} ({tb.LP.iloc[0]:.0f} LP, level with {tb.index[1]}, ahead on aggregate).")
    if g:
        ref, refname = R[(p, False)][team], f"{team} without Proposal 2"
    elif p == "unchanged" or REF_TEAM[team] is None:
        ref, refname = None, None
    else:
        ref, refname = R[("unchanged", False)][REF_TEAM[team]], f"{REF_TEAM[team]} (unchanged format)"
    if ref is not None:
        rank = f"Overall rank vs {refname}: {pos_changes(tb.Pos, ref['table'].Pos)}."
    else:
        rank = "Overall rank: baseline." if p == "unchanged" else "Overall rank: new table, no baseline."
    if not g:
        gen = "Gender shortfall: n/a (Proposal 2 off)."
    else:
        base, ev = R[(p, False)][team], []
        lk = r["lacked"] & r["count"].gt(0)
        for club in lk.index:
            for leg in lk.columns:
                if lk.loc[club, leg]:
                    s0, s1 = base["score"].loc[club, leg], r["score"].loc[club, leg]
                    r0, r1 = base["rank"].loc[club, leg], r["rank"].loc[club, leg]
                    flag = ", LOWER" if r1 > r0 else ""
                    ev.append(f"{club} L{leg} ({int(r['count'].loc[club, leg])}/{r['n']} archers, "
                              f"{(s1 - s0) / s0 * 100:+.0f}% score, rank {r0:.0f}->{r1:.0f}{flag})")
        gen = ("Gender shortfall (team not full): " + "; ".join(ev) + ".") if ev else "Gender shortfall: none."
    out = [lead, rank, gen]
    ties = [f"{c} L{l}" for c, l in r["unresolved_ties"]]
    if ties:
        out.append("Unresolved score ties: " + ", ".join(ties) + ".")
    return "\n".join(out)


def write_table(ws, r0, c0, title, tab, rank, legs_counted, counts=None, sub=None):
    """Write a league table (counts None) or participation table at (r0, c0).

    Rows: title, optional subtitle (sub, str), header, one row per club (in league order of tab).
    """
    ws.cell(r0, c0, title).font = Font(bold=True, size=12)
    if sub:
        ws.cell(r0 + 1, c0, sub).font = Font(italic=True, size=10)
        r0 += 1
    heads = ["Pos", "Club"] + [f"L{l}" for l in LEGS] + (["LP", "Agg"] if counts is None else ["Total", ""])
    for j, h in enumerate(heads):
        c = ws.cell(r0 + 1, c0 + j, h or None)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD; c.alignment = Alignment(horizontal="center")
    for i, club in enumerate(tab.index):
        vals = [i + 1, club]
        for l in LEGS:
            if l not in legs_counted:
                vals.append("n/c")
            elif counts is None:
                v = rank.loc[club, l]
                vals.append("-" if pd.isna(v) else int(v))
            else:
                vals.append(int(counts.loc[club, l]))
        vals += [int(tab.loc[club, "LP"]), int(tab.loc[club, "Agg"])] if counts is None else [int(counts.loc[club].sum()), None]
        for j, v in enumerate(vals):
            c = ws.cell(r0 + 2 + i, c0 + j, v)
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if j == 1 else "center")
            if counts is None and j in (0, 1, 10, 11):
                c.fill = SUB


def set_widths(ws, c0):
    """Column widths for the three 13-column tables of one block starting at column c0."""
    for t in range(3):
        s = c0 + t * TABLE_STRIDE
        ws.column_dimensions[col(s)].width = 5
        ws.column_dimensions[col(s + 1)].width = 13
        for k in range(2, 11):
            ws.column_dimensions[col(s + k)].width = 5
        ws.column_dimensions[col(s + 11)].width = 9
        ws.column_dimensions[col(s + 12)].width = 9


def build_explanation(wb):
    """Create the stand-alone 'Explanation' sheet (formats, rules, assumptions, data) for readers without the repo."""
    ws = wb.create_sheet("Explanation", 0)
    ws.column_dimensions["A"].width = 34
    for c, w in zip("BCDE", (14, 40, 50, 50)):
        ws.column_dimensions[c].width = w
    ws.sheet_view.showGridLines = False
    row = [1]

    def title(text):
        ws.cell(row[0], 1, text).font = Font(bold=True, size=16)
        row[0] += 2

    def head(text):
        for c in range(1, 6):
            ws.cell(row[0], c).fill = HEAD
        ws.cell(row[0], 1, text).font = Font(bold=True, size=12, color="FFFFFF")
        row[0] += 1

    def para(text, lines=1):
        c = ws.cell(row[0], 1, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row[0], start_column=1, end_row=row[0], end_column=5)
        ws.row_dimensions[row[0]].height = 15 * lines
        row[0] += 1

    def table(header, rows, heights, span_from=None):
        """Bordered table; if span_from is set, the cell in that column is merged across to column E."""
        for j, h in enumerate(header):
            c = ws.cell(row[0], 1 + j, h)
            c.font = Font(bold=True); c.fill = SUB; c.border = BOX
        row[0] += 1
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = ws.cell(row[0], 1 + j, v)
                c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
            if span_from:
                ws.merge_cells(start_row=row[0], start_column=span_from, end_row=row[0], end_column=5)
            ws.row_dimensions[row[0]].height = heights[i]
            row[0] += 1
        row[0] += 1

    title("SWWU 2025-26: proposed team formats - explanation")
    head("Purpose")
    para("The South Western and Welsh Universities (SWWU) archery league scores clubs on team results in each of 8 legs. "
         "This workbook re-scores the 2025-26 results under proposed changes to the team formats and shows how the league "
         "tables, and the number of archers counting towards each team, would change. The 'Team format analysis' sheet holds "
         "the results.", 3)
    row[0] += 1
    head("Data and scope")
    para("Clubs (8): Aberystwyth, Bath, Bristol, Cardiff, Exeter, Plymouth, Swansea, UWE. Teams are made up of archers from the "
         "same club. Results are the 2025-26 archer scores for legs 1-8; archers not in the cleaned results are excluded, "
         "alumni are not in any team, and archers who did not shoot (score 0) are not counted. Gender uses the Archery GB "
         "categories Open and Female, as recorded for each archer in each leg. Novices are archers in the novice classes.", 3)
    row[0] += 1
    table(["Leg", "Setting", "Round", "Notes", ""], [
        ["1, 3, 5", "Indoor", "Portsmouth", "60 arrows", ""], ["2", "Indoor", "WA 18m", "60 arrows", ""],
        ["4", "Indoor", "Worcester", "60 arrows", ""], ["6", "Indoor", "Bray 1", "30 arrows", ""],
        ["7", "Outdoor", "WA 900", "90 arrows", ""],
        ["8", "Outdoor", "WA 720 family",
         "Archers shoot different distances/faces (WA 70m, WA 60m, WA 50m 122cm, WA 50m 80cm). Scores are converted with the "
         "Archery GB handicap system to an equivalent score at a common round so a team of archers shooting different rounds "
         "can be compared: WA 70m for ENCT/ST, WA 60m for NNCT, WA 50m 122cm for UST. The compound team (CT) is not "
         "converted (all compound archers shot the same round).", ""]],
        [18, 18, 18, 18, 18, 80], span_from=4)
    head("Current team format (the 'Unchanged format' rows)")
    table(["Team", "Size", "Who can score", "Gender / other", ""], [
        ["ENCT - Experienced Non-Compound Team", "Best 4", "Non-compound archers (recurve, barebow, longbow); novices included",
         "No gender requirement", ""],
        ["NNCT - Novice Non-Compound Team", "Best 3", "Non-compound novices only", "No gender requirement", ""],
        ["CT - Compound Team", "Best 2", "Compound archers; novice and experienced together", "No gender requirement", ""]],
        [32, 32, 32])
    head("Proposals")
    table(["Proposal", "What changes", "", "", ""], [
        ["Proposal 1", "The compound team is removed and an Unsighted Team (UST) is added: best 3 barebow and longbow archers "
         "(novices can be included). The ENCT is reduced to 3 archers (recurve, barebow and longbow only; no compound) so the total number of "
         "archers across all teams stays the same. Compound archers are not eligible for any team.", "", "", ""],
        ["Proposal 1a", "As Proposal 1, but the ENCT is renamed Sighted Team (ST), takes recurve archers plus at most 1 compound "
         "archer, and compound archers may only score in the indoor legs (legs 1-6). Barebow/longbow archers are not eligible "
         "for the ST. Compound archers are not eligible for the novice team. Legs 7 and 8 do not count for the league "
         "(shown 'n/c').", "", "", ""],
        ["Proposal 2 (independent of 1 / 1a)", "Every team must include both gender categories (at least 1 Female and 1 Open "
         "archer). Where a club's eligible archers for that team do not include both categories, only (team size - 1) archers "
         "contribute to the team score. For the current-format compound team this means one Female and one Open archer.",
         "", "", ""]],
        [48, 62, 48], span_from=2)
    head("Common rules and assumptions")
    for t, n in [
        ("- An archer may be in more than one team if they meet that team's requirements (for example, a novice can score in "
         "both the ENCT/ST and the NNCT).", 1),
        ("- Incomplete teams are allowed: any club with at least one eligible archer in a leg is ranked. Under Proposal 2 the "
         "'team size - 1' cap only applies when the club's eligible archers lack Female or Open.", 2),
        ("- Team score = sum of the scores of the counted archers (the best total, subject to the team's rules). Under Proposal 2, "
         "if both genders are available, the best team containing at least one of each is chosen.", 2),
        ("- League points per leg: 10 for the highest team score, 9 for second, and so on; clubs with no team score 0. League "
         "ties (equal points) are broken on aggregate (sum of team scores).", 2),
        ("- Equal team scores in a leg are broken by the round's rule, summed over the team's archers: Portsmouth and Bray 1 - "
         "hits then golds; Worcester - hits, golds (5s), Xs; WA 18m - 10s then 9s; WA 900 and WA 720 - golds (10+X) then Xs. "
         "Where scores remain equal the clubs share the better rank (this did not occur).", 3),
        ("- Gender is taken as recorded per archer per leg. Leg 8 team scores use the handicap-equivalent scores described above.", 1)]:
        para(t, n)
    row[0] += 1
    head("How to read the 'Team format analysis' sheet")
    para("Rows are the three formats (Unchanged, Proposal 1, Proposal 1a); columns are the gender rule (Proposal 2 off on the left, "
         "on the right), giving 6 combinations. Each combination has three league tables side by side, with a participation "
         "table beneath each and an impact comment under that. " + HOW_TO_READ.replace("How to read: ", ""), 8)
    ws.sheet_view.zoomScale = 90


def main():
    R = run_all()
    wb = Workbook(); ws = wb.active; ws.title = "Team format analysis"
    ws.sheet_view.zoomScale = 60
    ws.cell(1, 1, "SWWU 2025-26: impact of proposed team formats").font = Font(bold=True, size=14)
    ws.cell(2, 1, HOW_TO_READ).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=4, end_column=41)
    for r in (2, 3, 4):
        ws.row_dimensions[r].height = 30
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 2
    for c in (42, 43, 44):
        ws.column_dimensions[col(c)].width = 2
    scale = dict(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84",
                 end_type="max", end_color="63BE7B")
    for gi, g in enumerate((False, True)):
        c_t = 3 + gi * BLOCK_OFFSET
        set_widths(ws, c_t)
        c = ws.cell(5, c_t, COL_HEADINGS[g])
        c.font = Font(bold=True, size=13, color="FFFFFF"); c.fill = HEAD; c.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=5, start_column=c_t, end_row=5, end_column=c_t + 3 * TABLE_STRIDE - 2)
    for bi, p in enumerate(PROPOSALS):
        top = BAND0 + bi * BAND
        h = ws.cell(top, 1, ROW_HEADINGS[p])
        h.font = Font(bold=True, size=13, color="FFFFFF"); h.fill = HEAD
        h.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        ws.merge_cells(start_row=top, start_column=1, end_row=top + 28, end_column=1)
        for gi, g in enumerate((False, True)):
            for ti, (t, r) in enumerate(R[(p, g)].items()):
                c0 = 3 + gi * BLOCK_OFFSET + ti * TABLE_STRIDE
                sub = SUBTITLE[p][t] + (GENDER_SUB if g else "")
                write_table(ws, top, c0, f"{t} league table", r["table"], r["rank"], r["legs"], sub=sub)
                write_table(ws, top + 12, c0, f"{t} participation", r["table"], None, r["legs"], counts=r["count"])
                ws.conditional_formatting.add(
                    f"{col(c0 + 2)}{top + 14}:{col(c0 + 1 + max(r['legs']))}{top + 21}", ColorScaleRule(**scale))
                cell = ws.cell(top + 23, c0, "Impact: " + table_comment(p, g, t, R))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.merge_cells(start_row=top + 23, start_column=c0, end_row=top + 28, end_column=c0 + 11)
    build_explanation(wb)
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
